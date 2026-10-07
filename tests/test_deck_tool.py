from pptx import Presentation
from pptx.util import Inches
import fill_deck

def test_fill_deck_writes_only_given_values(tmp_path):
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_table(2, 6, Inches(1), Inches(1), Inches(8), Inches(1)).table
    for i, h in enumerate(["Task", "Precision", "Recall", "False Positive", "False Negative", "Fix"]): t.cell(0, i).text = h
    t.cell(1, 0).text = "Classroom Headcount"
    for i in range(1, 5): t.cell(1, i).text = "[ ]"
    s.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text_frame.text = "[measure on test footage]"
    s.shapes.add_textbox(0, Inches(2), Inches(3), Inches(1)).text_frame.text = "CPU load: [measure]."
    src, dst = str(tmp_path / "a.pptx"), str(tmp_path / "b.pptx"); prs.save(src)
    acc = tmp_path / "acc.csv"; acc.write_text("Task,Precision,Recall,False Positive,False Negative\nClassroom Headcount,0.9,0.8,2,3\n")
    done = fill_deck.fill(src, dst, str(acc), {"headcount_mae": 1.2}, None, None)
    out = Presentation(dst); cells = [c.text for c in out.slides[0].shapes[0].table.rows[1].cells]
    assert cells[1:5] == ["0.9", "0.8", "2", "3"]
    texts = [sh.text_frame.text for sh in out.slides[0].shapes if sh.has_text_frame]
    assert "1.2 people average error" in texts and "CPU load: [measure]." in texts      # CPU stays blank: no speed file given
