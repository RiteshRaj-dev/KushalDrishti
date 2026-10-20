import os
import cv2, numpy as np, pytest
pytest.importorskip("pymupdf")
from pptx import Presentation
from pptx.util import Inches
import make_screenshots, insert_screenshots
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def test_detection_frame_notice_png_and_insert(tmp_path):
    vid = str(tmp_path / "v.mp4"); w = cv2.VideoWriter(vid, cv2.VideoWriter_fourcc(*"mp4v"), 10, (320, 180))
    for _ in range(30): w.write(np.random.default_rng(1).integers(0, 255, (180, 320, 3), dtype=np.uint8))
    w.release()
    p2 = make_screenshots.detection_frame(vid, str(tmp_path / "s2.png"))
    p3 = make_screenshots.notice_png(os.path.join(ROOT, "docs", "sample_notice_draft.pdf"), str(tmp_path / "s3.png"))
    assert cv2.imread(p2).shape[1] == 320 and cv2.imread(p3).shape[0] > 500
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6])
    for k in (1, 2, 3):
        s.shapes.add_shape(1, Inches(k), Inches(1), Inches(1.4), Inches(0.8)).name = f"Screenshot_{k}_placeholder"
    src, dst = str(tmp_path / "a.pptx"), str(tmp_path / "b.pptx"); prs.save(src)
    assert insert_screenshots.insert(src, dst, [p2, p3, p3]) == [1, 2, 3]
    names = [sh.name for sh in Presentation(dst).slides[0].shapes]
    assert names.count("Screenshot_1") == 1 and "Screenshot_1_placeholder" not in names
