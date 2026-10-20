"""Puts the screenshots into the dashed frames of the deck (shapes named Screenshot_1_placeholder, _2_, _3_).
    python tools/insert_screenshots.py SIH26245.pptx SIH26245_shots.pptx shots/shot1_map.png shots/shot2_detection.png shots/shot3_notice.png"""
import argparse, os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

def insert(src, dst, images):
    prs, placed = Presentation(src), []
    for slide in prs.slides:
        for k, img in enumerate(images, 1):
            if not img or not os.path.exists(img):
                continue
            holes = [s for s in slide.shapes if s.name == f"Screenshot_{k}_placeholder"]
            if not holes:
                continue
            h = holes[0]; x, y, w, ht = h.left, h.top, h.width, h.height
            for s in list(slide.shapes):               # remove the hint icon of this frame
                if s.name == f"Screenshot_{k}_icon":
                    s._element.getparent().remove(s._element)
            pic = slide.shapes.add_picture(img, x, y)
            box, im = w / ht, pic.width / pic.height   # crop to fill the frame without stretching
            if im > box:
                pic.crop_left = pic.crop_right = (1 - box / im) / 2
            else:
                pic.crop_top = pic.crop_bottom = (1 - im / box) / 2
            pic.left, pic.top, pic.width, pic.height = x, y, w, ht
            try:
                pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
            except Exception:
                pass
            pic.line.color.rgb = RGBColor.from_string("1565C0"); pic.line.width = Emu(12700)
            pic.name = f"Screenshot_{k}"
            h._element.getparent().remove(h._element)
            placed.append(k)
    prs.save(dst)
    return placed

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("src"); ap.add_argument("dst"); ap.add_argument("images", nargs="*")
    a = ap.parse_args()
    print("Placed screenshots:", insert(a.src, a.dst, a.images) or "none (frames not found?)")
